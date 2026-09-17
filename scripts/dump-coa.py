import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb["Setup"]
rows = list(ws.iter_rows(min_row=3, values_only=True))
print("--- ผังบัญชี (รหัส | หมวดบัญชี | กลุ่ม | บรรทัด P&L) ---")
for r in rows:
    code = r[5] if len(r) > 5 else None
    if code and str(code).strip() and str(code).strip()[0].isdigit():
        acct = str(r[6]).replace("\n", " ")[:52] if r[6] else ""
        grp = str(r[7]).replace("\n", " ")[:28] if r[7] else ""
        pl = str(r[8]).replace("\n", " ")[:34] if len(r) > 8 and r[8] else ""
        print(f"  {str(code).strip():<6} | {acct:<52} | {grp:<28} | {pl}")
print("\n--- กลุ่มเจ้าของ / เจ้าของ / บัญชี ---")
for idx, label in [(10, "Entity"), (11, "Owner"), (12, "Bank/Channel")]:
    seen = []
    for r in rows:
        v = r[idx] if len(r) > idx else None
        if v and str(v).strip() and str(v).strip() not in seen:
            seen.append(str(v).strip())
    print(f"  {label}: " + " · ".join(seen[:16]))
wb.close()
