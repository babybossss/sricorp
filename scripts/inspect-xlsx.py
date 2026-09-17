"""อ่านโครงสร้างไฟล์ Excel แล้วสรุป — ห้ามโหลดทั้งไฟล์เข้า context (CLAUDE.md กฎข้อ 6)"""
import sys
from openpyxl import load_workbook

for path in sys.argv[1:]:
    wb = load_workbook(path, read_only=True, data_only=True)
    print("=" * 70)
    print("FILE:", path.split("/")[-1])
    for ws in wb.worksheets:
        rows = ws.max_row or 0
        cols = ws.max_column or 0
        print(f"\n  SHEET: {ws.title}  ({rows} แถว x {cols} คอลัมน์)")
        it = ws.iter_rows(min_row=1, max_row=min(rows, 4), values_only=True)
        for i, row in enumerate(it):
            cells = ["" if c is None else str(c)[:28] for c in row[:14]]
            if not any(cells):
                continue
            print(f"    r{i+1}: " + " | ".join(cells))
    wb.close()
