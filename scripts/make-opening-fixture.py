"""
สร้างไฟล์ตัวอย่างขนาดเล็กสำหรับเทสต์ตัวอ่านไฟล์ยอดตั้งต้น — จากเทมเพลตจริงของลูกพี่ (templates/SRI_OS_2_...)
ใช้เทมเพลตจริงเพื่อให้เทสต์ตรวจ "โครงไฟล์ที่ลูกพี่จะส่งมา" ไม่ใช่โครงที่เราเดาเอง

รัน: /usr/bin/python3 scripts/make-opening-fixture.py
ผลลัพธ์: src/lib/db/import/__tests__/fixtures/opening-valid.xlsx  (ตัวเลขสมมติ ไม่ใช่ข้อมูลจริง)
"""
import sys
from datetime import date
from openpyxl import load_workbook

SRC = "templates/SRI_OS_2_ยอดตั้งต้นงบดุล.xlsx"
OUT = "src/lib/db/import/__tests__/fixtures/opening-valid.xlsx"

wb = load_workbook(SRC)

# owner sheet → { code: amount }, cash split [(account name, amount)]
DATA = {
    "BS-SRI Corporation": (
        {"1100": 2_000_000, "1400": 1_000_000, "1500": 5_000_000, "2410": 3_000_000, "3100": 4_000_000, "3300": 1_000_000},
        [("SRI - SCB", 2_000_000)],
    ),
    "BS-ธนากร": (
        {"1100": 1_000_000, "1200": 20_000, "1310": 100_000, "1500": 3_000_000, "1700": 300_000, "2410": 500_000,
         "3100": 2_000_000, "3300": 1_920_000},
        [("ธนากร - BBL 888", 600_000), ("ธนากร - SCB", 400_000)],
    ),
    "BS-สุธี": (
        {"1100": 300_000, "2310": 100_000, "3100": 200_000},
        [("สุธี - TTB", 300_000)],
    ),
}

for ws in wb.worksheets:
    if ws.title not in DATA:
        continue
    codes, banks = DATA[ws.title]
    for r in range(5, ws.max_row + 1):
        c = ws.cell(row=r, column=1).value
        if isinstance(c, str) and c in codes:
            ws.cell(row=r, column=5).value = codes[c]
    # ตารางแยกเงินสด: ระหว่างแถวหัว (B = 'บัญชี') กับ 'รวมเงินสด'
    head = next(r for r in range(1, ws.max_row + 1) if ws.cell(row=r, column=2).value == "บัญชี")
    end = next(r for r in range(head, ws.max_row + 1) if ws.cell(row=r, column=2).value == "รวมเงินสด")
    slots = list(range(head + 1, end))
    for i, (name, amount) in enumerate(banks):
        ws.cell(row=slots[i], column=2).value = name
        ws.cell(row=slots[i], column=5).value = amount
    for r in slots[len(banks):]:
        ws.cell(row=r, column=2).value = None
        ws.cell(row=r, column=3).value = None
        ws.cell(row=r, column=4).value = None
        ws.cell(row=r, column=5).value = None

# ชีทรายล็อต (ธนากร 300,000 = 70,000 + 100,000 + 129,900 + ค่าธรรมเนียม 100)
lots = wb["พอร์ตลงทุน (รายล็อต)"]
rows = [
    ("ธนากร", "หุ้นไทย", "PTT", date(2025, 3, 3), 2000, 35, None, "THB"),
    ("ธนากร", "ETF", "SET50", date(2025, 6, 2), 100, 1000, None, "THB"),
    ("ธนากร", "กองทุน", "K-FIXED", date(2025, 9, 1), 1299, 100, 100, "THB"),
]
for i, (owner, kind, name, d, units, price, fee, cur) in enumerate(rows):
    r = 6 + i
    lots.cell(row=r, column=1).value = owner
    lots.cell(row=r, column=2).value = kind
    lots.cell(row=r, column=3).value = name
    lots.cell(row=r, column=4).value = d  # วันที่จริงของ Excel (ตัวเลข serial) — ตัวอ่านต้องรับได้
    lots.cell(row=r, column=5).value = units
    lots.cell(row=r, column=6).value = price
    if fee is not None:
        lots.cell(row=r, column=7).value = fee
    lots.cell(row=r, column=9).value = cur

wb.save(OUT)
print("wrote", OUT)
