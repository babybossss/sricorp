import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1], read_only=True, data_only=True)
ws = wb["Setup"]
rows = list(ws.iter_rows(min_row=2, values_only=True))
def col(idx, label):
    print(f"\n--- {label} ---")
    seen = []
    for r in rows:
        v = r[idx] if idx < len(r) else None
        if v and str(v).strip() and str(v).strip() not in seen:
            seen.append(str(v).strip())
    for s in seen:
        print(" ", s.replace("\n", " / "))
    return seen
t = col(0, "ประเภทรายการ (Transaction Type)")
print("\n--- ทิศทาง+CF+PL ต่อประเภท ---")
for r in rows[1:]:
    if r[0] and str(r[0]).strip():
        print(f"  {str(r[0]).strip():<45} | dir={r[1]} | cf={r[2]} | pl={r[3]}")
wb.close()
